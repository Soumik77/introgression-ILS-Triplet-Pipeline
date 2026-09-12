from __future__ import annotations

import csv
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
RUN = BASE / "outputs" / "supervisor_scope" / "supervisor_scope_corrected_deadline_normalized"
REPORT = RUN / "comparison" / "final_report"
OUT = BASE / "deliverables"
ASSETS = OUT / "paper_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

DOCX_OUT = OUT / "Learning_Rooted_Triplets_Conference_Paper.docx"

BODY_FONT = "Times New Roman"
INK = "111111"
ACCENT = "1F4E79"
LIGHT_ACCENT = "D9EAF7"
MID_GRAY = "707070"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def make_figures() -> dict[str, Path]:
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Nimbus Roman", "Times New Roman", "DejaVu Serif"],
            "font.size": 7.0,
            "axes.titlesize": 8.0,
            "axes.labelsize": 7.5,
            "xtick.labelsize": 6.5,
            "ytick.labelsize": 6.5,
            "legend.fontsize": 6.2,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )

    paths: dict[str, Path] = {}

    # Pipeline overview.
    fig, ax = plt.subplots(figsize=(3.35, 2.08), dpi=300)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    boxes = [
        (0.06, 0.72, "Rooted backbone\n4-180 taxa"),
        (0.55, 0.72, "NMSC gene trees\nILS + 0/1 introgression"),
        (0.06, 0.37, "Enhanced17 triplets\n3 topology classes"),
        (0.55, 0.37, "Grouped CV prediction\ncalibrated probabilities"),
        (0.305, 0.04, "Normalized FM assembly\nrooted tree + nRF"),
    ]
    for x, y, label in boxes:
        rect = plt.Rectangle(
            (x, y), 0.39, 0.19, facecolor="#eef4f8", edgecolor="#1f4e79", linewidth=0.8
        )
        ax.add_patch(rect)
        ax.text(x + 0.195, y + 0.095, label, ha="center", va="center", fontsize=6.6)
    arrow = dict(arrowstyle="->", color="#404040", lw=0.9, shrinkA=2, shrinkB=2)
    ax.annotate("", xy=(0.55, 0.815), xytext=(0.45, 0.815), arrowprops=arrow)
    ax.annotate("", xy=(0.255, 0.56), xytext=(0.255, 0.72), arrowprops=arrow)
    ax.annotate("", xy=(0.55, 0.465), xytext=(0.45, 0.465), arrowprops=arrow)
    ax.annotate("", xy=(0.50, 0.23), xytext=(0.745, 0.37), arrowprops=arrow)
    ax.text(0.5, 0.965, "Evaluation pipeline", ha="center", va="top", fontsize=8.3, weight="bold")
    path = ASSETS / "figure_pipeline.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["pipeline"] = path

    # Paper-ready version of the taxon-count performance figure.
    rows = read_csv(REPORT / "table_internal_by_taxon_count.csv")
    selected = {
        "Enhanced17 ML soft FM": ("#1f4e79", "o", "ML-soft"),
        "Frequency FM": ("#b45f06", "s", "Frequency"),
        "Majority FM": ("#4f7f3b", "^", "Majority"),
    }
    fig, axes = plt.subplots(2, 1, figsize=(3.35, 3.55), dpi=300, sharex=True)
    for method, (color, marker, label) in selected.items():
        group = [r for r in rows if r["method"] == method]
        group.sort(key=lambda r: int(r["tree_size"]))
        x = np.array([int(r["tree_size"]) for r in group])
        exact = np.array([float(r["exact_match_accuracy"]) for r in group])
        nrf = np.array([float(r["mean_normalized_rooted_rf"]) for r in group])
        axes[0].plot(x, exact, color=color, lw=1.05, marker=marker, markevery=10, ms=2.2, label=label)
        axes[1].plot(x, nrf, color=color, lw=1.05, marker=marker, markevery=10, ms=2.2, label=label)
    for ax in axes:
        ax.axvline(20, color="#777777", ls="--", lw=0.75)
        ax.axvline(40, color="#999999", ls=":", lw=0.75)
        ax.set_ylim(-0.015, 1.02)
        ax.grid(True, color="#dddddd", lw=0.35, alpha=0.7)
    axes[0].set_ylabel("Exact-match accuracy")
    axes[0].set_title("A. Exact recovery", pad=2)
    axes[0].legend(frameon=False, ncol=3, loc="upper right", handlelength=1.4, columnspacing=0.7)
    axes[1].set_ylabel("Mean normalized rooted RF")
    axes[1].set_xlabel("Number of taxa")
    axes[1].set_title("B. Topological error", pad=2)
    axes[1].set_xlim(4, 180)
    axes[1].set_xticks([4, 20, 40, 80, 120, 160, 180])
    fig.tight_layout(pad=0.55)
    path = ASSETS / "figure_internal_by_taxon_count_paper.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["taxon"] = path

    # Model-versus-oracle figure.
    oracle_rows = read_csv(REPORT / "table_ml_soft_vs_oracle_by_taxon_stratum.csv")
    labels = [r["taxon_group"] for r in oracle_rows]
    model_exact = np.array([float(r["exact_match_accuracy_model"]) for r in oracle_rows])
    oracle_exact = np.array([float(r["exact_match_accuracy_oracle"]) for r in oracle_rows])
    model_nrf = np.array([float(r["mean_normalized_rooted_rf_model"]) for r in oracle_rows])
    oracle_nrf = np.array([float(r["mean_normalized_rooted_rf_oracle"]) for r in oracle_rows])
    x = np.arange(len(labels))
    width = 0.36
    fig, axes = plt.subplots(2, 1, figsize=(3.35, 3.42), dpi=300, sharex=True)
    axes[0].bar(x - width / 2, model_exact, width, color="#1f4e79", label="OOF ML-soft")
    axes[0].bar(x + width / 2, oracle_exact, width, color="#9cc2e5", label="Oracle labels")
    axes[1].bar(x - width / 2, model_nrf, width, color="#1f4e79")
    axes[1].bar(x + width / 2, oracle_nrf, width, color="#9cc2e5")
    for ax in axes:
        ax.set_ylim(0, 1.03)
        ax.grid(True, axis="y", color="#dddddd", lw=0.35, alpha=0.7)
    axes[0].set_ylabel("Exact-match accuracy")
    axes[0].set_title("A. Exact recovery", pad=2)
    axes[0].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].set_ylabel("Mean normalized rooted RF")
    axes[1].set_title("B. Topological error", pad=2)
    axes[1].set_xticks(x, labels, rotation=0)
    axes[1].set_xlabel("Taxon-count stratum")
    fig.tight_layout(pad=0.55)
    path = ASSETS / "figure_model_vs_oracle_paper.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["oracle"] = path

    # Fixed-budget coverage curve.
    taxa = np.arange(4, 181)
    total = np.array([math.comb(int(n), 3) for n in taxa], dtype=float)
    retained = np.minimum(total, 1000.0)
    share = 100.0 * retained / total
    fig, ax = plt.subplots(figsize=(3.35, 2.05), dpi=300)
    ax.plot(taxa, share, color="#1f4e79", lw=1.4)
    ax.fill_between(taxa, share, 0.08, color="#d9eaf7", alpha=0.85)
    ax.set_yscale("log")
    ax.set_xlim(4, 180)
    ax.set_ylim(0.08, 120)
    ax.set_xlabel("Number of taxa")
    ax.set_ylabel("Retained triplets (%)")
    ax.grid(True, which="both", color="#dddddd", lw=0.35, alpha=0.75)
    for n, dx, dy in [(20, 7, 1.25), (40, 8, 1.35), (180, -37, 1.65)]:
        y = 100.0 * min(math.comb(n, 3), 1000) / math.comb(n, 3)
        ax.scatter([n], [y], s=11, color="#b45f06", zorder=3)
        ax.annotate(
            f"n={n}: {y:.3g}%",
            xy=(n, y),
            xytext=(n + dx, y * dy),
            fontsize=6.3,
            arrowprops=dict(arrowstyle="-", lw=0.45, color="#666666"),
        )
    fig.tight_layout(pad=0.55)
    path = ASSETS / "figure_triplet_coverage.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["coverage"] = path

    return paths


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=35, start=45, bottom=35, end=45) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        data = kwargs[edge]
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in data:
                element.set(qn(f"w:{key}"), str(data[key]))


def set_run_font(run, size: float = 9.0, bold: bool | None = None, italic: bool | None = None, color: str = INK) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.64)
    section.right_margin = Inches(0.64)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.13)
    normal.paragraph_format.widow_control = True

    for name, size in (("Heading 1", 10.8), ("Heading 2", 9.5)):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(5 if name == "Heading 1" else 3)
        style.paragraph_format.space_after = Pt(1.5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    caption = doc.styles.add_style("Paper Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = BODY_FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    caption.font.size = Pt(7.5)
    caption.paragraph_format.space_before = Pt(1)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.first_line_indent = Inches(0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption.paragraph_format.keep_together = True

    ref = doc.styles.add_style("Paper Reference", WD_STYLE_TYPE.PARAGRAPH)
    ref.font.name = BODY_FONT
    ref._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    ref.font.size = Pt(7.7)
    ref.paragraph_format.space_after = Pt(1.2)
    ref.paragraph_format.first_line_indent = Inches(-0.13)
    ref.paragraph_format.left_indent = Inches(0.13)
    ref.paragraph_format.line_spacing = 0.95
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.core_properties.title = "Learning Rooted Triplets for Species-Tree Inference under ILS and Introgression"
    doc.core_properties.author = "Soumik Datta"
    doc.core_properties.subject = "Conference paper working draft"
    doc.core_properties.keywords = "species tree, rooted triplets, incomplete lineage sorting, introgression, machine learning"


def set_columns(section, num: int, space_twips: int = 300) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def set_section_geometry(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.64)
    section.right_margin = Inches(0.64)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, 7.2, color=MID_GRAY)


def add_paragraph(doc: Document, text: str, *, indent: bool = True, bold_lead: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.13 if indent else 0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, 9, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, 9, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, 9, italic=italic)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    set_run_font(r, 10.8 if level == 1 else 9.5, bold=True)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, 9, italic=True)


def add_numbered_contribution(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(1.2)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p.add_run(f"{number}. ")
    set_run_font(r1, 9, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, 9)


def add_figure(doc: Document, path: Path, caption: str, width: float = 3.42) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(style="Paper Caption")
    cp.paragraph_format.keep_together = True
    label, rest = caption.split(" ", 1)
    r1 = cp.add_run(label + " ")
    set_run_font(r1, 7.5, bold=True)
    r2 = cp.add_run(rest)
    set_run_font(r2, 7.5)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    cp = doc.add_paragraph(style="Paper Caption")
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(1)
    cp.paragraph_format.keep_with_next = True
    label, rest = caption.split(" ", 1)
    r1 = cp.add_run(label + " ")
    set_run_font(r1, 7.5, bold=True)
    r2 = cp.add_run(rest)
    set_run_font(r2, 7.5)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    if widths:
        total_twips = int(sum(widths) * 1440)
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(total_twips))
        tbl_w.set(qn("w:type"), "dxa")
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        grid_cols = table._tbl.tblGrid.gridCol_lst
        for j, width in enumerate(widths):
            twips = str(int(width * 1440))
            if j < len(grid_cols):
                grid_cols[j].set(qn("w:w"), twips)
            table.columns[j].width = Inches(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_ACCENT)
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[j])
        p = cell.paragraphs[0]
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_run_font(r, 7.2, bold=True)
        set_cell_border(cell, top={"val": "single", "sz": 8, "color": ACCENT}, bottom={"val": "single", "sz": 5, "color": ACCENT})

    for i, row in enumerate(rows):
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for j, value in enumerate(row):
            cell = cells[j]
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, 7.1)
            if i == len(rows) - 1:
                set_cell_border(cell, bottom={"val": "single", "sz": 8, "color": ACCENT})
    tail = doc.add_paragraph()
    tail.paragraph_format.first_line_indent = Inches(0)
    tail.paragraph_format.space_after = Pt(1)
    tail.paragraph_format.line_spacing = 0.5


def add_reference(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Paper Reference")
    r = p.add_run(text)
    set_run_font(r, 7.7)


def build_document(figs: dict[str, Path]) -> Path:
    doc = Document()
    configure_doc(doc)
    set_columns(doc.sections[0], 1)
    add_page_number(doc.sections[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("Learning Rooted Triplets for Species-Tree Inference under ILS and Introgression:\nAccuracy, Assembly, and Sampling Limits")
    set_run_font(r, 15.2, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Soumik Datta")
    set_run_font(r, 9.5, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Department of Computer Science and Engineering\nBangladesh University of Engineering and Technology (BUET), Dhaka, Bangladesh")
    set_run_font(r, 8.3)

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_geometry(body_section)
    set_columns(body_section, 2, 270)
    body_section.footer.is_linked_to_previous = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(1.5)
    r = p.add_run("Abstract")
    set_run_font(r, 9.3, bold=True)

    abstract = (
        "Gene-tree discordance caused by incomplete lineage sorting (ILS) and introgression complicates species-tree inference. "
        "We study a pipeline that summarizes simulated gene-tree collections as rooted triplets, predicts one of three rooted topologies from 17 frequency and branch-length features, and assembles a complete rooted tree with a normalized Fiduccia-Mattheyses-style objective. "
        "The evaluation contains 5,310 simulations spanning every taxon count from 4 to 180, 1,770 held-out rooted backbones, three ILS levels, three introgression strengths, and 500 to 2,500 gene trees per simulation. "
        "The grouped out-of-fold classifier attains 94.93% triplet accuracy. In the diagnostic 4-20-taxon stratum, calibrated soft weighting recovers 273 of 510 trees exactly (53.53%), compared with 39.80% for frequency weighting and 42.55% for majority weighting. "
        "An oracle using true labels on the same sampled triplets recovers 500 of 510 trees (98.04%), showing that most remaining low-taxon error arises before assembly. By contrast, the oracle falls to 4.00% exact recovery at 21-40 taxa and zero above 40 taxa. "
        "The fixed 1,000-triplet cap covers 87.72% of possible triplets at 20 taxa but only 0.1046% at 180 taxa. These results separate three failure sources: an initially biased split objective, local classification error, and a dominant large-tree information deficit. "
        "The current method improves low-taxon reconstruction over internal controls, but larger trees require taxon-dependent, coverage-aware triplet selection and independent external benchmarking."
    )
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(abstract)
    set_run_font(r, 8.4)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(3)
    r1 = p.add_run("Keywords: ")
    set_run_font(r1, 8.1, bold=True)
    r2 = p.add_run("species-tree inference, rooted triplets, incomplete lineage sorting, introgression, gradient boosting, tree assembly")
    set_run_font(r2, 8.1, italic=True)

    add_heading(doc, "1 Introduction")
    add_paragraph(
        doc,
        "Species-tree inference becomes difficult when individual loci follow genealogies that disagree with one another and with the species history. ILS produces such discordance even without gene flow, and introgression adds a second mechanism by allowing some loci to follow a reticulate path (Maddison 1997; Degnan and Rosenberg 2009; Solis-Lemus and Ane 2016). Summary methods therefore seek a species-level history from collections of gene trees rather than treating every locus as if it shared one topology."
    )
    add_paragraph(
        doc,
        "Rooted triplets are the smallest rooted phylogenetic statements: for taxa a, b, and c, exactly three binary resolutions are possible. They are attractive because a compatible dense triplet set determines a rooted tree, and because triplet frequencies retain local evidence that can be aggregated by a scalable assembler (Aho et al. 1981; Liu, Yu, and Edwards 2010; Islam et al. 2020). Yet a local-to-global pipeline can fail at several distinct stages. A classifier can choose the wrong triplet, the assembly objective can prefer a structurally biased split, or the sampled triplets can leave much of a large tree unconstrained."
    )
    add_paragraph(
        doc,
        "This distinction matters for interpreting exact-match accuracy. A triplet classifier may be correct on most local decisions while the final tree is not exact, because exact recovery requires every nontrivial rooted clade to agree with the reference. Conversely, a poor final tree does not by itself show that the classifier is poor. A diagnostic must hold the sampled triplet set and assembler fixed while replacing predictions with the true triplet labels."
    )
    add_paragraph(doc, "We address three research questions:", indent=False)
    add_numbered_contribution(doc, 1, "RQ1: Does a learned 17-feature triplet classifier improve complete-tree reconstruction over frequency and majority controls when coverage is dense?")
    add_numbered_contribution(doc, 2, "RQ2: How much error is attributable to triplet classification, the assembly score, and sampled-triplet coverage?")
    add_numbered_contribution(doc, 3, "RQ3: What happens when one fixed triplet budget is applied from 4 through 180 taxa?")
    add_paragraph(doc, "The study makes four contributions:", indent=False)
    add_numbered_contribution(doc, 1, "A leakage-controlled classifier using topology frequencies, conditional branch-length summaries, asymmetry features, and observed gene-tree count.")
    add_numbered_contribution(doc, 2, "A normalized FM-style split objective whose oracle behavior is checked against the exact rooted-tree evaluator.")
    add_numbered_contribution(doc, 3, "A controlled evaluation of 5,310 simulations covering 177 taxon counts and 1,770 backbone topologies under ILS and zero or one introgression edge.")
    add_numbered_contribution(doc, 4, "An oracle diagnosis showing that classifier error dominates in the dense low-taxon regime, whereas sparse triplet coverage dominates the large-tree stress test.")

    add_heading(doc, "2 Related Work")
    add_paragraph(
        doc,
        "Coalescent summary methods address gene-tree discordance without concatenating all loci. MP-EST estimates a rooted species tree by maximizing a pseudo-likelihood derived from rooted triplet distributions (Liu, Yu, and Edwards 2010). ASTRAL instead maximizes quartet agreement within a constrained search space and is statistically consistent under the multispecies coalescent (Mirarab et al. 2014). STELAR formalizes constrained triplet consensus and uses dynamic programming to maximize agreement with triplets induced by rooted gene trees (Islam et al. 2020). These methods establish triplets and quartets as useful local summaries, but their statistical guarantees assume particular data-generating models and do not directly answer how a learned triplet rule behaves under introgression."
    )
    add_paragraph(
        doc,
        "Triplet aggregation also appears in the supertree literature. SuperTriplets optimizes a triplet-based median criterion through agglomeration and local rearrangements (Ranwez, Criscuolo, and Douzery 2010). Triplet MaxCut recursively partitions taxa to maximize consistency with rooted triplets (Sevillya, Frenkel, and Snir 2016). Our assembler is related in spirit but uses three weights per sampled taxon triplet and a normalized FM-style local search. Because the external programs did not complete all selected simulations in the available run, we do not claim an external benchmark against them."
    )
    add_paragraph(
        doc,
        "Explicit network methods model both ILS and gene flow. SNaQ uses quartet concordance factors in a network pseudolikelihood, while PhyloNet provides likelihood, pseudo-likelihood, and Bayesian approaches for reticulate histories (Solis-Lemus and Ane 2016; Wen et al. 2018). Our target is narrower: the generating rooted backbone, not the complete reticulation network. The experiment asks whether local gene-tree summaries can recover the major vertical topology when the gene trees were generated on a network containing at most one introgression edge."
    )

    add_heading(doc, "3 Method")
    add_heading(doc, "3.1 Inference target and workflow", 2)
    add_paragraph(
        doc,
        "For a taxon set V and an unordered triplet {a,b,c}, sorted labels define class 0 as ((a,b),c), class 1 as ((a,c),b), and class 2 as ((b,c),a). The target class is induced by the generating rooted backbone. Predictor variables are calculated only from the simulated gene trees. Gamma, ILS level, donor, recipient, pair class, backbone identity, and the generating tree are retained as metadata and excluded from the classifier. Figure 1 summarizes the workflow."
    )
    add_figure(
        doc,
        figs["pipeline"],
        "Figure 1. Rooted backbones generate network-coalescent gene trees. Triplet features feed grouped out-of-fold prediction, and the weighted triplets are assembled and evaluated as complete rooted trees.",
    )

    add_heading(doc, "3.2 Simulation design", 2)
    add_paragraph(
        doc,
        "The design includes every integer taxon count from 4 through 180. Ten unique rooted binary backbones were selected at each size across four normalized Colless imbalance strata, giving 1,770 backbones. Internal branch lengths were 0.02, 0.10, or 0.50 coalescent units. Introgression strength gamma was 0, 0.25, or 0.50. For positive gamma, donor-recipient pairs represented close, intermediate, and distant backbone separations. Requested gene-tree counts were 500, 1,000, 2,000, or 2,500."
    )
    add_table(
        doc,
        "Table 1. Executed simulation and learning design.",
        ["Element", "Executed specification"],
        [
            ["Taxa", "4-180; 177 sizes"],
            ["Backbones", "10 per size; 1,770 total"],
            ["ILS lengths", "0.02, 0.10, 0.50 coalescent units"],
            ["Gamma", "0, 0.25, 0.50"],
            ["Gene trees", "500, 1,000, 2,000, 2,500"],
            ["Candidate grid", "191,160 simulations"],
            ["Executed set", "3 per backbone; 5,310 total"],
            ["Generated trees", "7,976,000 gene trees"],
            ["Held-out rows", "4,975,320 triplet rows"],
        ],
        widths=[1.02, 2.25],
    )
    add_paragraph(
        doc,
        "The full factorial grid contained 191,160 candidate simulations and 286,740,000 candidate gene trees. A deterministic factor-balanced subset selected three cells per backbone before outcomes were observed. Every backbone contributed one simulation at each gamma level and one at each ILS level; gene-tree counts and positive pair classes differed in frequency by at most one within each taxon size. This reduced the run to 5,310 simulations and 7,976,000 gene trees while retaining all sizes and backbones. Gene trees were generated with PhyloCoalSimulations under a network multispecies coalescent model (Fogg, Allman, and Ane 2023). No sequence evolution or gene-tree estimation error was included."
    )

    add_heading(doc, "3.3 Triplet sampling and Enhanced17", 2)
    add_paragraph(
        doc,
        "All C(n,3) triplets were retained when this number was at most 1,000, which is exhaustive through 19 taxa. From 20 taxa onward, a deterministic sampler first ensured that every taxon appeared and then added unique random triplets until reaching 1,000. The same identifiers were used for feature extraction, prediction, and assembly within a simulation."
    )
    add_paragraph(
        doc,
        "For candidate topology k, let n_k be its support among G gene trees. Its frequency is p_k = n_k/G. For the two alternatives i and j, an asymmetry feature was computed as |p_i-p_j|/(p_i+p_j), with zero assigned when the denominator was zero. Enhanced17 contains three topology frequencies; the majority-topology asymmetry; the mean, median, and population standard deviation of induced internal branch length for each topology; three candidate-specific asymmetries; and log(G). Missing conditional branch summaries were median-imputed inside the fitted pipeline with missingness indicators."
    )

    add_heading(doc, "3.4 Grouped learning and calibration", 2)
    add_paragraph(
        doc,
        "The classifier was scikit-learn's HistGradientBoostingClassifier with learning rate 0.06, 250 maximum iterations, 31 maximum leaf nodes, L2 regularization 1.0, early stopping, and seed 20260828 (Pedregosa et al. 2011). Each model used at most 20 sampled rows per training simulation to bound memory, but prediction covered every retained triplet in the test simulations. Cyclic topology-coordinate augmentation was applied only to the training partition."
    )
    add_paragraph(
        doc,
        "Five outer folds were grouped by backbone and stratified by taxon count. Because each size had ten backbones, each fold held out two backbones at every size: 354 unseen backbones and 1,062 simulations. Three grouped inner folds estimated one temperature parameter from 161 values between 0.25 and 4.0 by minimizing multiclass log loss. Test labels were never used for fitting or calibration."
    )

    add_heading(doc, "3.5 Normalized FM-style assembly", 2)
    add_paragraph(
        doc,
        "For each selected triplet, the assembler retained three topology weights. At a candidate split L|R, a topology contributed positively when its paired taxa remained together, negatively when that pair was separated, and zero when all three taxa fell on one side. Only unique triplets crossing the split entered the objective. The corrected score was the signed support sum divided by the number of unique crossing triplet identifiers."
    )
    add_equation(doc, "score(L|R) = [sum over crossing triplets of signed support] / [number of unique crossing triplets].")
    add_paragraph(
        doc,
        "The search followed an FM-style move-and-lock heuristic (Fiduccia and Mattheyses 1982). It initialized a bipartition from a low-support seed pair, assigned remaining taxa by pair affinity, moved the unlocked taxon with the largest score gain, and retained the best intermediate partition. Recursion continued until each terminal subset contained one taxon. The normalization is essential: without it, a split can gain score merely by resolving more triplets."
    )

    add_heading(doc, "3.6 Evaluation and oracle diagnosis", 2)
    add_paragraph(
        doc,
        "We compare calibrated probability weights (ML-soft), a one-hot predicted class (ML-hard), winner-confidence weights (ML-confidence), observed topology frequencies (Frequency), and the most frequent topology as a one-hot vector (Majority). All methods use the same normalized assembler and simulation identifiers."
    )
    add_paragraph(
        doc,
        "Inferred and reference trees were represented by their nontrivial rooted clades. Rooted RF distance is the clade-set symmetric difference; normalized rooted RF divides it by the total number of clades in both trees (Robinson and Foulds 1981). Exact match is true precisely when normalized rooted RF equals zero. Exact accuracy uses all requested simulations in its denominator. We report paired mean differences, 95% percentile intervals from 10,000 bootstrap resamples clustered by backbone, and one-sided Wilcoxon signed-rank tests on backbone means with Holm correction over four baselines and two metrics."
    )
    add_paragraph(
        doc,
        "The oracle replaces every predicted triplet label with its stored true backbone label while preserving the exact sampled triplet set, assembler, taxon set, and simulation. It is therefore a conditional ceiling for the implemented assembly under the selected triplets, not a ceiling for all possible triplets."
    )

    add_heading(doc, "4 Results")
    add_heading(doc, "4.1 Triplet prediction", 2)
    add_paragraph(
        doc,
        "Across 4,975,320 held-out triplet rows, Enhanced17 obtained 94.9345% accuracy, multiclass log loss 0.1248, multiclass Brier score 0.0705, and 10-bin top-label expected calibration error 0.0021. These are outer-fold results. The model training sample contained 105,420 rows before coordinate augmentation and 316,260 rows after augmentation; total recorded training wall time was 573.4 seconds."
    )

    add_heading(doc, "4.2 The assembly correction", 2)
    add_paragraph(
        doc,
        "The initial unnormalized split sum failed an oracle test. With true triplet labels for all 300 simulations from 11 through 20 taxa, it recovered 166 trees exactly (55.33%) and had mean normalized rooted RF 0.0438. Dividing by the number of unique crossing triplets increased recovery to 290 of 300 (96.67%) and reduced mean normalized rooted RF to 0.0022. Every size from 11 through 19 then reached 30 of 30 oracle exact matches. At 20 taxa, the oracle recovered 20 of 30 because the cap retained 1,000 of 1,140 possible triplets. This test identifies the original low-taxon deficit as an assembly-objective bug rather than a weakness of true triplet labels."
    )

    add_heading(doc, "4.3 Internal comparison at 4-20 taxa", 2)
    add_paragraph(
        doc,
        "The 4-20-taxon diagnostic stratum contains 510 simulations from 170 held-out backbones. It was defined during post-run diagnosis, so its inferential statistics are exploratory rather than prespecified."
    )
    add_table(
        doc,
        "Table 2. Complete-tree performance at 4-20 taxa (510 simulations).",
        ["Method", "Exact trees (%)", "Mean nRF"],
        [
            ["ML-soft", "273 (53.53)", "0.1781"],
            ["ML-hard", "273 (53.53)", "0.1781"],
            ["ML-confidence", "268 (52.55)", "0.1803"],
            ["Frequency", "203 (39.80)", "0.2480"],
            ["Majority", "217 (42.55)", "0.2325"],
        ],
        widths=[1.25, 1.23, 0.78],
    )
    add_paragraph(
        doc,
        "ML-soft improved exact accuracy over Frequency by 0.1373 (95% clustered bootstrap interval 0.1078 to 0.1667; Holm-adjusted p = 8.59 x 10^-13) and reduced normalized rooted RF by 0.0700 (0.0537 to 0.0870; adjusted p = 9.92 x 10^-14). Against Majority, its exact gain was 0.1098 (0.0804 to 0.1392; adjusted p = 4.62 x 10^-10) and its normalized RF reduction was 0.0544 (0.0373 to 0.0723; adjusted p = 9.39 x 10^-8)."
    )
    add_paragraph(
        doc,
        "The three ML weighting rules were not distinguishable. ML-soft and ML-hard recovered the same 273 trees, with exact gain 0 and adjusted p = 1.000. ML-soft exceeded ML-confidence by 0.0098 exact accuracy, but the interval included zero and adjusted p was 0.675. The measurable improvement therefore comes from the learned local decision rule, not from a demonstrated advantage of calibrated soft weights over hard or winner-confidence weights."
    )

    add_heading(doc, "4.4 Accuracy changes with tree size", 2)
    add_figure(
        doc,
        figs["taxon"],
        "Figure 2. Exact-match accuracy and mean normalized rooted RF by taxon count. Vertical lines mark 20 and 40 taxa. All methods use the fixed 1,000-triplet cap from 20 taxa onward.",
    )
    add_paragraph(
        doc,
        "Figure 2 shows a sharp size transition. ML-soft exact accuracy is 60.00% over 4-10 taxa and 49.00% over 11-20 taxa, then falls to 1.83% over 21-40 taxa and zero above 40. Mean normalized rooted RF increases from 0.1602 over 11-20 taxa to 0.3981 over 21-40, 0.7998 over 41-80, and 0.9722 over 121-180. Frequency and Majority follow the same scaling pattern, which argues against a failure unique to probability weighting."
    )

    add_heading(doc, "4.5 Oracle separates model and coverage", 2)
    add_table(
        doc,
        "Table 3. ML-soft and conditional oracle exact recovery by taxon stratum.",
        ["Taxa", "n", "ML-soft", "Oracle"],
        [
            ["4-10", "210", "60.00%", "100.00%"],
            ["11-20", "300", "49.00%", "96.67%"],
            ["21-40", "600", "1.83%", "4.00%"],
            ["41-80", "1,200", "0.00%", "0.00%"],
            ["81-120", "1,200", "0.00%", "0.00%"],
            ["121-180", "1,800", "0.00%", "0.00%"],
        ],
        widths=[0.78, 0.62, 0.93, 0.93],
    )
    add_figure(
        doc,
        figs["oracle"],
        "Figure 3. Out-of-fold ML-soft results and the conditional oracle using true labels for the same sampled triplets. The oracle removes classification error but does not add unselected triplets.",
    )
    add_paragraph(
        doc,
        "At 4-20 taxa, the oracle recovered 500 of 510 trees (98.04%), compared with 273 for ML-soft. Dense correct triplets therefore make the normalized assembler almost exact, and classification remains the principal low-taxon gap. At 21-40 taxa, however, even the oracle recovered only 24 of 600 trees. It recovered none of 4,200 trees above 40 taxa. The large-tree failure remains after classification error is removed and must be attributed to the information supplied to assembly or to assembly behavior under that sparse information."
    )

    add_heading(doc, "4.6 Fixed-budget sparsity", 2)
    add_figure(
        doc,
        figs["coverage"],
        "Figure 4. Fraction of all possible rooted taxon triplets retained by the fixed budget. The y-axis is logarithmic; all triplets are used through 19 taxa.",
    )
    add_paragraph(
        doc,
        "The combinatorial explanation is direct. At 20 taxa, 1,000 selected triplets cover 87.72% of the 1,140 possible triplets. At 40 taxa, they cover 10.12% of 9,880; at 80 taxa, 1.22% of 82,160; and at 180 taxa, 0.1046% of 955,860. Pair coverage also becomes impossible to guarantee: 1,000 triplets contain at most 3,000 taxon-pair incidences, whereas 180 taxa contain 16,110 distinct pairs. Even with no repeated pair incidence, the upper bound is 18.62%. The stress test therefore reduces information as tree size grows."
    )

    add_heading(doc, "4.7 Full-scope result and runtime", 2)
    add_paragraph(
        doc,
        "Across all 5,310 simulations, ML-soft recovered 284 trees (5.3484%) with mean normalized rooted RF 0.7836. ML-hard recovered 282 (5.3107%), ML-confidence 281 (5.2919%), Frequency 213 (4.0113%), and Majority 225 (4.2373%). Every method returned a fully resolved tree for every simulation; there were no evaluation failures or taxon mismatches, and exact match was equivalent to normalized rooted RF equal to zero in every row."
    )
    add_paragraph(
        doc,
        "These pooled values are dominated by the 4,200 simulations above 40 taxa and should not be used alone to summarize practical accuracy. Mean assembly time per tree ranged from 0.101 seconds for ML-hard to 0.169 seconds for Frequency, excluding simulation, feature extraction, and model training. Runtime was not the limiting factor in the implemented assembly; information coverage was."
    )

    add_heading(doc, "5 Discussion")
    add_heading(doc, "5.1 Why exact accuracy is much lower than triplet accuracy", 2)
    add_paragraph(
        doc,
        "Triplet accuracy and exact-tree accuracy ask different questions. The classifier score averages millions of local three-taxon decisions. Exact match is a conjunction over the final tree's nontrivial clades: one incompatible high-level split makes the entire tree non-exact. Errors are also dependent, because many triplets support the same clade and the recursive assembler commits to early partitions. A 5.07% local error rate can therefore be concentrated on structurally important triplets rather than spread harmlessly across the tree."
    )
    add_paragraph(
        doc,
        "The oracle makes this explanation testable. In the dense 4-20 regime, replacing predictions with truth raises exact recovery from 53.53% to 98.04%. This is the classifier-and-weighting gap conditional on the current triplets and assembler. In the large-tree regime, replacing predictions with truth barely helps because almost all possible triplets were never supplied. The same low pooled accuracy thus contains two qualitatively different problems."
    )

    add_heading(doc, "5.2 What the method currently establishes", 2)
    add_paragraph(
        doc,
        "The supported positive claim is local and conditional. Enhanced17 improves complete-tree recovery over Frequency and Majority controls on the exploratory 4-20-taxon subset, and its probabilities are well calibrated at triplet level. The normalized split objective passes an exhaustive-or-nearly-exhaustive oracle check through 20 taxa. These results justify continued development of the learned triplet representation."
    )
    add_paragraph(
        doc,
        "The study does not establish accurate reconstruction through 180 taxa. Beyond 40 taxa, model and oracle both fail under the same fixed 1,000-triplet budget. Those results quantify a sampling design limit, not an upper bound on the method with a sufficient taxon-dependent budget. Likewise, soft, hard, and confidence weighting should be treated as equivalent on the present evidence."
    )

    add_heading(doc, "5.3 Design implications", 2)
    add_paragraph(
        doc,
        "The next assembler evaluation should vary information, not only optimization. Triplet selection should be deterministic and coverage-aware, ensuring repeated representation of each taxon and taxon pair. The cap should grow with taxon count or with a target pair-coverage criterion. A small independent sensitivity design can compare 1,000, 2,500, 5,000, and 10,000 triplets where feasible, with exhaustive sampling at smaller sizes. Oracle runs should precede classifier runs at every budget so that a cap incapable of recovering the true tree is rejected before expensive retraining."
    )
    add_paragraph(
        doc,
        "Classifier improvements should then target the residual low-taxon gap. Useful diagnostics include triplet accuracy by ILS, gamma, gene-tree count, pair class, and whether an error supports a high-level backbone split. Adding taxon count and retained-triplet proportion may make the model size-aware, but any model selection and calibration must remain inside grouped training folds. Multiple deterministic FM starts can test local-optimum sensitivity, although no optimizer can restore triplets that were never sampled."
    )

    add_heading(doc, "6 Limitations")
    add_paragraph(
        doc,
        "The executed design samples three of 108 candidate condition cells per backbone. It balances marginal factors but cannot estimate every within-backbone factorial interaction. The simulator uses one sampled lineage per taxon, zero or one introgression edge, common internal branch length within a simulation, and true gene trees. Sequence evolution, alignment error, and gene-tree estimation error are absent, so empirical accuracy may be lower. The target is the backbone tree, not the location or direction of the reticulation."
    )
    add_paragraph(
        doc,
        "The 4-20 and larger taxon strata were chosen after inspecting the initial run and must be replicated on new backbones before confirmatory claims. External STELAR, MP-EST, and Triplet MaxCut outputs were incomplete for the 5,310 selected simulations and were excluded. Consequently, the reported significance tests compare internal weighting controls only. There is no real-data analysis, and the runtime summary does not include the full simulation-to-feature pipeline."
    )

    add_heading(doc, "7 Reproducibility and Artifact Scope")
    add_paragraph(
        doc,
        "The implementation uses fixed seed 20260828 for design, simulation, training, and deterministic tie-breaking. Simulation identifiers, fold assignments, model metrics, per-method evaluation tables, paired comparisons, and the oracle outputs are retained in the accompanying internal results archive. The archive contains source, configuration, and final aggregate results, while large out-of-fold prediction shards and the training sample are excluded from the compact transfer package. A public repository and independent validation split should be prepared before submission."
    )

    add_heading(doc, "8 Conclusion")
    add_paragraph(
        doc,
        "A high local classification score did not translate directly into high exact-tree accuracy, but the reason changes with tree size. After correcting the split normalization, true triplet labels recovered 96.67% of 11-20-taxon trees and 98.04% across 4-20 taxa. In that dense regime, Enhanced17 improved exact recovery over Frequency and Majority controls, while the three ML weighting variants were statistically indistinguishable. Above 40 taxa, both model and oracle failed because a fixed 1,000-triplet cap represented a vanishing fraction of the available structure. The next technical priority is therefore taxon-dependent, pair-coverage-aware triplet selection, followed by independent internal replication and complete external benchmarking."
    )

    add_heading(doc, "References")
    references = [
        "Aho, A. V.; Sagiv, Y.; Szymanski, T. G.; and Ullman, J. D. 1981. Inferring a tree from lowest common ancestors with an application to the optimization of relational expressions. SIAM Journal on Computing, 10(3): 405-421. doi:10.1137/0210030.",
        "Degnan, J. H.; and Rosenberg, N. A. 2009. Gene tree discordance, phylogenetic inference and the multispecies coalescent. Trends in Ecology & Evolution, 24(6): 332-340. doi:10.1016/j.tree.2009.01.009.",
        "Fiduccia, C. M.; and Mattheyses, R. M. 1982. A linear-time heuristic for improving network partitions. In Proceedings of the 19th Design Automation Conference, 175-181.",
        "Fogg, J.; Allman, E. S.; and Ane, C. 2023. PhyloCoalSimulations: A simulator for network multispecies coalescent models, including a new extension for the inheritance of gene flow. Systematic Biology, 72(5): 1171-1179. doi:10.1093/sysbio/syad030.",
        "Islam, M.; Sarker, K.; Das, T.; Reaz, R.; and Bayzid, M. S. 2020. STELAR: a statistically consistent coalescent-based species tree estimation method by maximizing triplet consistency. BMC Genomics, 21: 136. doi:10.1186/s12864-020-6519-y.",
        "Liu, L.; Yu, L.; and Edwards, S. V. 2010. A maximum pseudo-likelihood approach for estimating species trees under the coalescent model. BMC Evolutionary Biology, 10: 302. doi:10.1186/1471-2148-10-302.",
        "Maddison, W. P. 1997. Gene trees in species trees. Systematic Biology, 46(3): 523-536.",
        "Mirarab, S.; Reaz, R.; Bayzid, M. S.; Zimmermann, T.; Swenson, M. S.; and Warnow, T. 2014. ASTRAL: genome-scale coalescent-based species tree estimation. Bioinformatics, 30(17): i541-i548. doi:10.1093/bioinformatics/btu462.",
        "Pedregosa, F.; Varoquaux, G.; Gramfort, A.; Michel, V.; Thirion, B.; Grisel, O.; Blondel, M.; Prettenhofer, P.; Weiss, R.; Dubourg, V.; Vanderplas, J.; Passos, A.; Cournapeau, D.; Brucher, M.; Perrot, M.; and Duchesnay, E. 2011. Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12: 2825-2830.",
        "Ranwez, V.; Criscuolo, A.; and Douzery, E. J. P. 2010. SuperTriplets: a triplet-based supertree approach to phylogenomics. Bioinformatics, 26(12): i115-i123. doi:10.1093/bioinformatics/btq196.",
        "Robinson, D. F.; and Foulds, L. R. 1981. Comparison of phylogenetic trees. Mathematical Biosciences, 53(1-2): 131-147. doi:10.1016/0025-5564(81)90043-2.",
        "Sevillya, G.; Frenkel, Z.; and Snir, S. 2016. Triplet MaxCut: a new toolkit for rooted supertree. Methods in Ecology and Evolution, 7(11): 1359-1365. doi:10.1111/2041-210X.12606.",
        "Solis-Lemus, C.; and Ane, C. 2016. Inferring phylogenetic networks with maximum pseudolikelihood under incomplete lineage sorting. PLOS Genetics, 12(3): e1005896. doi:10.1371/journal.pgen.1005896.",
        "Wen, D.; Yu, Y.; Zhu, J.; and Nakhleh, L. 2018. Inferring phylogenetic networks using PhyloNet. Systematic Biology, 67(4): 735-740. doi:10.1093/sysbio/syy015.",
    ]
    for ref in references:
        add_reference(doc, ref)

    add_heading(doc, "Appendix A: Low-Taxon Results by Size")
    add_paragraph(
        doc,
        "Table A1 reports the per-size diagnostic behind the pooled 4-20-taxon result. Each row contains 30 simulations from ten held-out backbones. Oracle recovery is complete through 19 taxa and falls only when the fixed cap first becomes non-exhaustive at 20 taxa.",
    )
    detailed = read_csv(
        RUN
        / "diagnostics"
        / "oracle_fm_4_180_normalized"
        / "oracle_vs_model_by_taxon_count.csv"
    )
    detailed_rows = []
    for row in detailed:
        n_taxa = int(row["tree_size"])
        if n_taxa > 20:
            break
        detailed_rows.append(
            [
                str(n_taxa),
                f'{100 * float(row["model_exact_accuracy"]):.1f}%',
                f'{100 * float(row["oracle_exact_accuracy"]):.1f}%',
                f'{float(row["model_mean_normalized_rooted_rf"]):.3f}',
                f'{float(row["oracle_mean_normalized_rooted_rf"]):.3f}',
            ]
        )
    add_table(
        doc,
        "Table A1. ML-soft and oracle performance for each taxon count from 4 through 20.",
        ["Taxa", "ML exact", "Oracle", "ML nRF", "Or. nRF"],
        detailed_rows,
        widths=[0.47, 0.77, 0.77, 0.63, 0.62],
    )
    add_paragraph(
        doc,
        "The per-size model curve is not monotonic within the dense regime because each size contains only 30 simulations with different backbones and balanced condition assignments. The pooled strata are therefore more stable summaries than any single size. The oracle pattern is qualitatively different: the same assembler is exact for every simulation through 19 taxa, then loses 10 trees at the first capped size.",
    )

    add_heading(doc, "Appendix B: Full-Scope Method Summary")
    add_table(
        doc,
        "Table B1. Corrected results over all 5,310 simulations.",
        ["Method", "Exact", "Mean", "Median", "Mean s"],
        [
            ["ML-soft", "284", "0.784", "0.927", "0.123"],
            ["ML-hard", "282", "0.783", "0.927", "0.101"],
            ["ML-conf.", "281", "0.784", "0.928", "0.123"],
            ["Frequency", "213", "0.806", "0.944", "0.169"],
            ["Majority", "225", "0.790", "0.928", "0.114"],
        ],
        widths=[0.92, 0.57, 0.59, 0.62, 0.56],
    )
    add_paragraph(
        doc,
        "Mean and median columns report normalized rooted RF; Mean s reports assembly seconds per tree. Because 79.1% of the simulations contain more than 40 taxa, these pooled values primarily reflect the deliberately sparse large-tree stress test.",
    )

    add_heading(doc, "Appendix C: Reporting Checklist")
    add_numbered_contribution(doc, 1, "Exact matching compares rooted topology only. Newick child order and branch lengths are ignored, and exact match is equivalent to normalized rooted RF equal to zero.")
    add_numbered_contribution(doc, 2, "Every method was requested and evaluated on 5,310 simulations, with zero failures, zero taxon mismatches, and fully resolved inferred trees.")
    add_numbered_contribution(doc, 3, "The 4-20-taxon comparison is an exploratory diagnostic selected after the initial run. Its p values require independent replication.")
    add_numbered_contribution(doc, 4, "Results above 40 taxa describe a fixed-budget stress test. They do not estimate accuracy under adequate triplet coverage.")
    add_numbered_contribution(doc, 5, "STELAR, MP-EST, and Triplet MaxCut were not included because complete matched outputs were unavailable for all 5,310 simulations.")

    # A final continuous one-column section makes Word balance the last
    # two-column reference page instead of leaving the second column empty.
    balance_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_geometry(balance_section)
    set_columns(balance_section, 1)
    balance_section.footer.is_linked_to_previous = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.first_line_indent = Inches(0)
    spacer.add_run("")

    # Reinforce consistent page geometry across sections. The footer was
    # created once on the first section and is linked through the document.
    for section in doc.sections:
        set_section_geometry(section)

    doc.save(DOCX_OUT)
    return DOCX_OUT


if __name__ == "__main__":
    figure_paths = make_figures()
    output = build_document(figure_paths)
    print(output)
